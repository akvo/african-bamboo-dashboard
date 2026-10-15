import datetime
import zipfile

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from api.v1.v1_a2f.models import A2FSnapshot

# Header row of each data sheet, names and order exactly as in the
# client's workbook. README and code_lists sheets are not read.
SHEETS = {
    "cooperatives": [
        "coop_id", "coop_name", "coop_type", "region", "woreda",
        "registration_no", "created_at", "updated_at",
    ],
    "members": [
        "member_id", "coop_id", "target_group", "first_name",
        "last_name", "gender", "date_of_birth", "region", "woreda",
        "kebele", "gps_latitude", "gps_longitude", "land_size_ha",
        "bamboo_area_ha", "income_baseline", "income_current",
        "training_completed", "status", "enrollment_date",
        "jobs_created",
    ],
    "loans": [
        "loan_id", "member_id", "cooperative_id", "loan_type",
        "loan_amount_etb", "interest_rate", "repayment_freq",
        "loan_status", "outstanding_balance", "par_30_flag",
        "green_flag", "collateral_type", "credit_score",
        "disbursement_date", "maturity_date",
    ],
    "parcels": [
        "parcel_id", "coop_id", "member_id", "title_deed_no",
        "title_deed_type", "centroid_lat", "centroid_long",
        "polygon_geojson", "area_ha", "soil_type", "soil_ph", "slope",
        "forest_cover_pct", "ndvi_score", "carbon_stock_tco2e",
        "land_use_type", "deforestation_risk", "annual_rainfall_mm",
        "elevation_m", "verified_by", "survey_date",
    ],
}
DATE_COLUMNS = {
    "created_at", "updated_at", "date_of_birth", "enrollment_date",
    "disbursement_date", "maturity_date", "survey_date",
}
NUMBER_COLUMNS = {
    "gps_latitude", "gps_longitude", "land_size_ha", "bamboo_area_ha",
    "income_baseline", "income_current", "training_completed",
    "jobs_created", "loan_amount_etb", "interest_rate",
    "outstanding_balance", "par_30_flag", "green_flag", "credit_score",
    "centroid_lat", "centroid_long", "area_ha", "soil_ph", "slope",
    "forest_cover_pct", "ndvi_score", "carbon_stock_tco2e",
    "annual_rainfall_mm", "elevation_m",
}


def is_blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def parse_date(value):
    if isinstance(value, datetime.datetime):
        return value.date().isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    return datetime.date.fromisoformat(str(value).strip()).isoformat()


def parse_number(value):
    if isinstance(value, bool):
        raise ValueError(value)
    if isinstance(value, (int, float)):
        return value
    number = float(str(value).strip())
    return int(number) if number.is_integer() else number


def parse_cell(column, value):
    if is_blank(value):
        return None
    if column in DATE_COLUMNS:
        return parse_date(value)
    if column in NUMBER_COLUMNS:
        return parse_number(value)
    return str(value).strip()


def read_sheet(wb, name, columns):
    if name not in wb.sheetnames:
        raise CommandError(f"Missing sheet '{name}'")
    rows = wb[name].iter_rows(values_only=True)
    header = list(next(rows, ()))
    while header and is_blank(header[-1]):
        header.pop()
    header = [str(h).strip() for h in header]
    if header != columns:
        raise CommandError(
            f"Sheet '{name}': header must be {columns}, got {header}"
        )
    records = []
    for number, row in enumerate(rows, start=2):
        row = row[: len(columns)]
        if all(is_blank(v) for v in row):
            continue
        record = {}
        for column, value in zip(columns, row):
            try:
                record[column] = parse_cell(column, value)
            except (TypeError, ValueError):
                raise CommandError(
                    f"Sheet '{name}' row {number} column "
                    f"'{column}': cannot read {value!r}"
                )
        records.append(record)
    return records


def read_parameters(wb):
    if "parameters" not in wb.sheetnames:
        raise CommandError("Missing sheet 'parameters'")
    raw = {}
    for row in wb["parameters"].iter_rows(min_row=2, values_only=True):
        if row and not is_blank(row[0]):
            raw[str(row[0]).strip()] = row[1] if len(row) > 1 else None
    try:
        data_as_of = parse_date(raw["data_as_of"])
    except (KeyError, AttributeError, TypeError, ValueError):
        raise CommandError(
            "Sheet 'parameters': data_as_of must be a YYYY-MM-DD date"
        )
    params = {
        k: None if is_blank(v) else str(v).strip()
        for k, v in raw.items()
    }
    params["data_as_of"] = data_as_of
    return params


class Command(BaseCommand):
    help = (
        "Replace all access to finance data with the workbook at PATH. "
        "Any file name works. In a deployed environment copy the file "
        "into the backend pod (kubectl cp) and run this there."
    )

    def add_arguments(self, parser):
        parser.add_argument("path")

    def handle(self, *args, **options):
        # A file object, not a path: openpyxl refuses paths whose
        # extension it does not recognise, and the name must not matter.
        try:
            with open(options["path"], "rb") as f:
                wb = load_workbook(f, read_only=True, data_only=True)
                sheets = {
                    name: read_sheet(wb, name, columns)
                    for name, columns in SHEETS.items()
                }
                params = read_parameters(wb)
        except (OSError, zipfile.BadZipFile, InvalidFileException) as e:
            raise CommandError(f"Cannot open {options['path']}: {e}")
        with transaction.atomic():
            A2FSnapshot.objects.all().delete()
            A2FSnapshot.objects.create(
                data_as_of=params["data_as_of"],
                parameters=params,
                **sheets,
            )
        counts = ", ".join(f"{len(v)} {k}" for k, v in sheets.items())
        self.stdout.write(
            f"Imported {counts} as of {params['data_as_of']}"
        )
