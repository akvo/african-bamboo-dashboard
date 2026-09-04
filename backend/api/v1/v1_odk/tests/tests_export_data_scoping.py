"""TC02 data correctness across all three formats.

Two claims the QA case makes about every export:

1. With the Approved filter active, only approved
   records appear.
2. Values reflect the latest cleaned data, not the
   original uncleaned submission.

Both already hold, but nothing pinned them. These
tests run generate_export_file end to end (the task
function directly -- the queue is not synchronous in
tests) and read the produced files back.
"""

import json
import os
import tempfile
import zipfile

import shapefile as shp
from openpyxl import load_workbook

from django.test import TestCase, override_settings

from api.v1.v1_jobs.constants import JobStatus, JobTypes
from api.v1.v1_jobs.models import Jobs
from api.v1.v1_odk.constants import ApprovalStatusTypes
from api.v1.v1_odk.models import (
    FarmerFieldMapping,
    FormMetadata,
    FormQuestion,
    Plot,
    Submission,
)
from api.v1.v1_odk.tasks import generate_export_file
from api.v1.v1_odk.tests.mixins import OdkTestHelperMixin
from utils.storage import get_path

VALID_WKT = (
    "POLYGON(("
    "38.47 7.05,"
    "38.48 7.05,"
    "38.48 7.06,"
    "38.47 7.06,"
    "38.47 7.05))"
)


@override_settings(USE_TZ=False, TEST_ENV=True)
class ExportDataScopingTest(
    TestCase, OdkTestHelperMixin
):
    def setUp(self):
        self.user = self.create_kobo_user()
        self.auth = self.get_auth_header()
        self.form = FormMetadata.objects.create(
            asset_uid="scope-form-1",
            name="Scope Test Form",
            polygon_field="geoshape",
            region_field="region",
            sub_region_field="woreda",
            plot_name_field="farmer_name",
        )
        for qname, label in [
            ("farmer_name", "Farmer Name"),
            ("region", "Region"),
        ]:
            FormQuestion.objects.create(
                form=self.form,
                name=qname,
                label=label,
                type="text",
            )
        # The xlsx Plot sheet carries only IDs and
        # geometry; cleaned text values surface in the
        # Farmer sheet, which is driven by this mapping.
        FarmerFieldMapping.objects.create(
            form=self.form,
            unique_fields="farmer_name",
            values_fields="farmer_name,region",
        )
        self.approved = self._create_plot(
            "Approved Farmer",
            ApprovalStatusTypes.APPROVED,
        )
        self._create_plot("Pending Farmer", None)
        self._create_plot(
            "Rejected Farmer",
            ApprovalStatusTypes.REJECTED,
        )
        self._cleanup = []
        self.addCleanup(self._remove_files)

    def _remove_files(self):
        for rel in self._cleanup:
            full = get_path(rel)
            if os.path.exists(full):
                os.remove(full)

    def _create_plot(self, name, approval_status):
        sub = Submission.objects.create(
            uuid=f"sub-scope-{Plot.objects.count()}",
            form=self.form,
            kobo_id=str(200 + Plot.objects.count()),
            submission_time=1700000000000,
            raw_data={
                "geoshape": (
                    "7.05 38.47 0 0;"
                    "7.06 38.47 0 0;"
                    "7.06 38.48 0 0;"
                    "7.05 38.48 0 0;"
                    "7.05 38.47 0 0"
                ),
                "farmer_name": name,
                "region": "Amhara",
                "woreda": "Bahir Dar",
            },
            approval_status=approval_status,
        )
        return Plot.objects.create(
            submission=sub,
            form=self.form,
            plot_name=name,
            polygon_wkt=VALID_WKT,
            min_lat=7.05,
            max_lat=7.06,
            min_lon=38.47,
            max_lon=38.48,
            region="Amhara",
            sub_region="Bahir Dar",
            created_at=1700000000000,
        )

    def _run_export(self, job_type, status=None):
        filters = {}
        if status:
            filters["status"] = status
        job = Jobs.objects.create(
            type=job_type,
            status=JobStatus.pending,
            created_by=self.user,
            info={
                "form_id": self.form.asset_uid,
                "filters": filters,
            },
        )
        generate_export_file(job.pk)
        job.refresh_from_db()
        self.assertEqual(
            job.status,
            JobStatus.done,
            job.result,
        )
        self._cleanup.append(
            job.info["file_path"]
        )
        return job

    # --- Approved-only scoping ---

    def test_geojson_approved_scope(self):
        job = self._run_export(
            JobTypes.export_geojson, "approved"
        )

        with open(
            get_path(job.info["file_path"])
        ) as fh:
            data = json.load(fh)

        self.assertEqual(job.info["record_count"], 1)
        self.assertEqual(len(data["features"]), 1)
        props = data["features"][0]["properties"]
        self.assertEqual(
            props["VAL_STATUS"], "approved"
        )

    def test_shapefile_approved_scope(self):
        job = self._run_export(
            JobTypes.export_shapefile, "approved"
        )

        self.assertEqual(job.info["record_count"], 1)
        records = self._read_shapefile_records(job)
        self.assertEqual(len(records), 1)

    def test_xlsx_approved_scope(self):
        """The 'XLSX: all plots' comment is stale.

        The status filter is applied before the format
        branch, so it applies to xlsx too.
        """
        job = self._run_export(
            JobTypes.export_xlsx, "approved"
        )

        rows = self._read_plot_rows(job)
        self.assertEqual(len(rows), 1)

    def test_no_filter_exports_every_record(self):
        job = self._run_export(
            JobTypes.export_geojson
        )

        self.assertEqual(job.info["record_count"], 3)

    # --- Cleaned data round-trip ---

    def _clean_value(self, **fields):
        """Edit a submission the way the DCU does."""
        with self._no_kobo_sync():
            resp = self.client.patch(
                "/api/v1/odk/submissions/"
                f"{self.approved.submission.uuid}"
                "/edit_data/",
                {"fields": fields},
                content_type="application/json",
                **self.auth,
            )
        self.assertEqual(resp.status_code, 200)

    def test_geojson_reflects_cleaned_value(self):
        self._clean_value(region="Oromia")
        job = self._run_export(
            JobTypes.export_geojson, "approved"
        )

        with open(
            get_path(job.info["file_path"])
        ) as fh:
            data = json.load(fh)

        props = data["features"][0]["properties"]
        self.assertEqual(props["REGION"], "Oromia")
        self.assertNotIn("Amhara", props.values())

    def test_shapefile_reflects_cleaned_value(self):
        self._clean_value(region="Oromia")
        job = self._run_export(
            JobTypes.export_shapefile, "approved"
        )

        records = self._read_shapefile_records(job)
        flat = [str(v) for v in records[0]]
        self.assertIn("Oromia", flat)
        self.assertNotIn("Amhara", flat)

    def test_xlsx_reflects_cleaned_value(self):
        """The Farmer sheet carries cleaned values."""
        self._clean_value(
            farmer_name="Corrected Name"
        )
        job = self._run_export(
            JobTypes.export_xlsx, "approved"
        )

        rows = self._read_farmer_rows(job)
        flat = [
            str(v) for r in rows for v in r if v
        ]
        self.assertIn("Corrected Name", flat)
        self.assertNotIn("Approved Farmer", flat)

    # --- helpers ---

    def _no_kobo_sync(self):
        from unittest.mock import patch

        return patch("api.v1.v1_odk.views.async_task")

    def _read_shapefile_records(self, job):
        full = get_path(job.info["file_path"])
        stem = os.path.basename(full)[: -len(".zip")]
        with tempfile.TemporaryDirectory() as tmpdir:
            with zipfile.ZipFile(full, "r") as zf:
                zf.extractall(tmpdir)
            reader = shp.Reader(
                os.path.join(tmpdir, stem)
            )
            return list(reader.records())

    def _read_farmer_rows(self, job):
        return self._read_rows(job, "Farmer table")

    def _read_plot_rows(self, job):
        return self._read_rows(job, "Plot Table")

    def _read_rows(self, job, sheet):
        full = get_path(job.info["file_path"])
        wb = load_workbook(full)
        ws = wb[sheet]
        return [
            row
            for row in ws.iter_rows(
                min_row=2, values_only=True
            )
            if any(v is not None for v in row)
        ]

    def test_farmer_sync_failure_does_not_fail_export(
        self,
    ):
        """Stale farmer data beats no export at all."""
        from unittest.mock import patch

        with patch(
            "api.v1.v1_odk.tasks"
            ".sync_farmers_for_form",
            side_effect=Exception("sync boom"),
        ):
            job = self._run_export(
                JobTypes.export_xlsx
            )

        self.assertEqual(job.status, JobStatus.done)
        self.assertIn(
            "sync boom",
            job.info["farmer_sync_error"],
        )

    def test_job_info_records_download_name(self):
        job = self._run_export(
            JobTypes.export_geojson
        )

        name = job.info["download_name"]
        self.assertTrue(
            name.startswith("Scope_Test_Form")
        )
        self.assertIn("scope-form-1", name)
        self.assertTrue(name.endswith(".geojson"))
