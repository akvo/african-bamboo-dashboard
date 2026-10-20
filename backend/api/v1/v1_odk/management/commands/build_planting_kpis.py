import datetime
import glob
import json
import os
import re
import zipfile
from collections import Counter

from django.core.management.base import BaseCommand, CommandError
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

# Rules from AB_KPI_Guidance_AKVO.docx. Header prefix in the planting
# lists (lowercased, files misspell some), Kobo column suffix in the
# survival export, display name.
SPECIES = [
    ("yushania", "ya", "Yushania alpina (highland bamboo)"),
    ("millet", "mf", "Millettia ferruginea"),
    ("hegenia", "ha", "Hagenia abyssinica"),
    ("mulberry", "mam", "Morus alba (mulberry)"),
]
PLANTING_SHEET = "seedling distrbution"
SURVIVAL_SHEET = "Seedling Survival Monitoring"
ER_SHEET = "Net GHG Benefit"


def number(value):
    if value is None or (isinstance(value, str) and not value.strip()):
        return 0
    return float(value)


def year_of(path):
    return int(re.search(r"_(\d{4})", os.path.basename(path)).group(1))


def latest(paths, pattern):
    if not paths:
        raise CommandError(f"No {pattern} file found")
    return max(paths, key=year_of)


def open_workbook(path):
    try:
        return load_workbook(path, read_only=True, data_only=True)
    except (OSError, zipfile.BadZipFile, InvalidFileException) as e:
        raise CommandError(f"Cannot open {path}: {e}")


def read_planting(path):
    """Plots, hectares and seedlings per species of one planting list.

    Without a plot size column the list's area comes from the
    ``_NNha`` part of its file name.
    """
    sheets = [
        ws
        for ws in open_workbook(path).worksheets
        if ws.title.lower().startswith(PLANTING_SHEET)
    ]
    if not sheets:
        raise CommandError(f"{path}: no 'Seedling distrbution' sheet")
    plots, area, has_plot_size = 0, 0.0, False
    seedlings = {name: 0 for _, _, name in SPECIES}
    for ws in sheets:
        rows = ws.iter_rows(values_only=True)
        header = [str(h or "").strip().lower() for h in next(rows, ())]
        columns = {}
        for prefix, _, name in SPECIES:
            match = [i for i, h in enumerate(header) if h.startswith(prefix)]
            if len(match) != 1:
                where = f"{path} '{ws.title}'"
                raise CommandError(f"{where}: no single '{prefix}' column")
            columns[name] = match[0]
        size = header.index("plot size") if "plot size" in header else None
        has_plot_size = has_plot_size or size is not None
        for row in rows:
            row = tuple(row) + (None,) * len(header)
            # A farmer row has a first name; the rest are blank padding.
            if row[0] is None or not str(row[0]).strip():
                continue
            plots += 1
            for name, i in columns.items():
                seedlings[name] += number(row[i])
            if size is not None:
                area += number(row[size])
    if not has_plot_size:
        match = re.search(r"_(\d+(?:\.\d+)?)ha", os.path.basename(path))
        if not match:
            raise CommandError(
                f"{path}: no plot size column and no _NNha in the file name"
            )
        area = float(match.group(1))
    return plots, area, seedlings


def farmer_key(value):
    """AB plus five digits; IDs are typed with spaces, O for 0,
    without the AB prefix or with a sixth digit. Cutting to the first
    five is the rule that reproduces the guidance's 264 plots, though
    it merges AB002551 into 255 (to confirm with the client)."""
    text = re.sub(r"\s+", "", str(value or "")).upper().replace("O", "0")
    digits = re.sub(r"\D", "", text)[:5]
    return f"AB{int(digits):05d}" if digits else text


def read_survival(path):
    ws = open_workbook(path)[SURVIVAL_SHEET]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    visits = {}
    for row in rows:
        visit = dict(zip(header, row))
        name = str(visit["farmer_name"] or "").strip().lower()
        if name == "test" or visit["date_monitoring"] is None:
            continue
        key = farmer_key(visit["farmer_id"])
        # Of two visits on the same day the first counts.
        if (
            key not in visits
            or visit["date_monitoring"] > visits[key]["date_monitoring"]
        ):
            visits[key] = visit

    def rate(suffix):
        kept = visits.values()
        planted = sum(number(v[f"total_planted_{suffix}"]) for v in kept)
        survived = sum(number(v[f"total_survived_{suffix}"]) for v in kept)
        return round(100 * survived / planted, 1)

    up_to = max(v["date_monitoring"] for v in visits.values())
    return {
        "rate": rate("all"),
        "plots": len(visits),
        "up_to": up_to.date().isoformat(),
        "species": {name: rate(suffix) for _, suffix, name in SPECIES},
    }


def read_carbon(path):
    """Total VCU issuance from Table 2 (vintage rows then a Total row
    in column L, issuance in column S)."""
    ws = open_workbook(path)[ER_SHEET]
    vintages = []
    for row in ws.iter_rows(min_col=12, max_col=19):
        label = str(row[0].value or "").strip()
        if label.lower() == "total" and vintages:
            start = vintages[0].split(" to ")[0]
            end = vintages[-1].split(" to ")[1]
            name = os.path.basename(path)
            return {
                "tco2e": round(row[7].value),
                "period_start": parse_day(start),
                "period_end": parse_day(end),
                "source": f"{name}, {ER_SHEET}!{row[7].coordinate}",
            }
        if " to " in label:
            vintages.append(label)
    raise CommandError(f"{path}: no Table 2 Total row in '{ER_SHEET}'")


def parse_day(text):
    return datetime.datetime.strptime(text, "%d-%b-%Y").date().isoformat()


class Command(BaseCommand):
    help = (
        "Print the Planting & carbon figures as JSON, rebuilt from the "
        "workbooks in DIR: every AB_Planting_<year>_<NN>ha.xlsx, the "
        "latest AB_Survival_Monitoring_<year>.xlsx and the latest "
        "AB_ER_Model_<year>.xlsx. Redirect the output to "
        "frontend/src/data/planting.json."
    )

    def add_arguments(self, parser):
        parser.add_argument("dir")

    def handle(self, *args, **options):
        def find(pattern):
            return glob.glob(os.path.join(options["dir"], pattern))

        planting = sorted(find("AB_Planting_*.xlsx"))
        if not planting:
            raise CommandError("No AB_Planting_*.xlsx file found")
        survival = read_survival(
            latest(find("AB_Survival_Monitoring_*.xlsx"), "survival")
        )
        carbon = read_carbon(latest(find("AB_ER_Model_*.xlsx"), "ER model"))

        plots, hectares, seedlings = 0, 0.0, Counter()
        for path in planting:
            p, area, counts = read_planting(path)
            plots += p
            hectares += area
            seedlings.update(counts)

        years = [year_of(p) for p in planting]
        data = {
            "planting_years": [min(years), max(years)],
            "hectares": round(hectares, 1),
            "plots": plots,
            "monitoring_up_to": survival["up_to"],
            "survival_rate": survival["rate"],
            "monitored_plots": survival["plots"],
            "species": [
                {
                    "name": name,
                    "seedlings": int(seedlings[name]),
                    "survival_rate": survival["species"][name],
                }
                for _, _, name in SPECIES
            ],
            "carbon": carbon,
        }
        self.stdout.write(json.dumps(data, indent=2))
